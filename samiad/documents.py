"""Fill the Word templates and turn them into PDFs (LibreOffice)."""
from __future__ import annotations

import re
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from docxtpl import DocxTemplate

from .rules import ADDONS, ALWAYS_SHOWN, Basis, Student, ZERO

TEMPLATES = Path(__file__).resolve().parent.parent / "templates"

CAMPUS_ADDRESSES = {
    "Eton College": "Eton, Windsor, Berkshire SL4 6DW",
    "Bromsgrove School": "Worcester Road, Bromsgrove, Worcestershire B61 7DU",
    "Trent College": "Derby Road, Long Eaton, Nottingham NG10 4AD",
    "RGS Surrey Hills": "Mickleham, Dorking, Surrey RH5 6EA",
    "Marlborough College": "Bath Road, Marlborough, Wiltshire SN8 1PA",
    "The Oratory School": "Woodcote, Reading RG8 0PJ",
    "The Leys School": "Trumpington Road, Cambridge CB2 7AD",
}


def gbp(x: Decimal) -> str:
    return f"{x:,.2f}"


def nice_date(value) -> str:
    """'2027-07-04' -> '4 July 2027'. HubSpot DOBs are text like '10/31/1976' (month first)."""
    if not value:
        return ""
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            d = datetime.strptime(s[:10] if fmt == "%Y-%m-%d" else s, fmt).date()
            return f"{d.day} {d:%B %Y}"
        except ValueError:
            continue
    return s


def parse_date(value) -> date | None:
    try:
        return date.fromisoformat(str(value)[:10]) if value else None
    except ValueError:
        return None


def clean_course(value: str | None) -> str:
    return re.sub(r"\s*\((SS|CS)\)\s*$", "", value or "").strip()


def nights(s: Student) -> int | None:
    a, d = parse_date(s.p("arrival_dats")), parse_date(s.p("departure_date"))
    return (d - a).days if a and d and d > a else None


def student_context(s: Student, basis: Basis) -> dict:
    course, spec = clean_course(s.p("course")), (s.p("specialism") or "").strip()
    programme = ", ".join(x for x in (course, spec) if x)
    n = nights(s)
    lines = [{"label": "Course Fee" if basis is Basis.GROSS else "Net Fee",
              "amount": f"£{gbp(s.course_price(basis))}"}]
    for prop, label, _ in ADDONS:
        amt = s.addon(prop)
        if amt != ZERO:
            lines.append({"label": label, "amount": f"£{gbp(amt)}"})
        elif prop == "airport_transfer_fee":
            lines.append({"label": label, "amount": "Included" if s.transfers_booked else "Not included"})
        elif prop in ALWAYS_SHOWN:
            lines.append({"label": label, "amount": "Not included"})
    receipt_desc = ", ".join(x for x in (s.p("campus"), spec, f"{n} nights" if n else "") if x)
    return {
        "first_name": s.p("firstname") or "", "last_name": s.p("lastname") or "",
        "date_of_birth": nice_date(s.p("date_of_birth")), "nationality": s.p("nationality") or "",
        "passport_number": s.p("passport_number") or "",
        "arrival_date": nice_date(s.p("arrival_dats")), "departure_date": nice_date(s.p("departure_date")),
        "nights": n or "", "campus": s.p("campus") or "",
        "campus_address": CAMPUS_ADDRESSES.get(s.p("campus") or "", ""),
        "course": programme, "specialism": "", "programme": programme,
        "transfers_booked": s.transfers_booked,
        "total": gbp(s.stated_total(basis)),
        "lines": lines,
        # receipt: first line describes the course
        "receipt_lines": [{"label": receipt_desc, "amount": lines[0]["amount"]}] + lines[1:],
    }


@dataclass
class BookingDocs:
    deal_id: str
    deal_name: str
    basis: Basis
    students: list[Student]
    customer_name: str
    owner_name: str
    owner_email: str
    total: Decimal
    received: Decimal = ZERO
    outstanding: Decimal = ZERO

    def context(self) -> dict:
        studs = [student_context(s, self.basis) for s in self.students]
        return {
            "deal_id": self.deal_id, "customer_name": self.customer_name,
            "owner_name": self.owner_name, "owner_email": self.owner_email,
            "today": nice_date(date.today().isoformat()),
            "total": gbp(self.total), "received": gbp(self.received),
            "outstanding": gbp(self.outstanding), "students": studs,
        }


def _render(template: str, ctx: dict) -> bytes:
    tpl = DocxTemplate(TEMPLATES / template)
    tpl.render(ctx, autoescape=True)
    buf = BytesIO()
    tpl.save(buf)
    return buf.getvalue()


def _suffix(basis: Basis) -> str:
    return "gross" if basis is Basis.GROSS else "net"


def confirmation_docx(b: BookingDocs) -> bytes:
    return _render(f"confirmation_{_suffix(b.basis)}.docx", b.context())


def receipt_docx(b: BookingDocs) -> bytes:
    ctx = b.context()
    for s in ctx["students"]:
        s["lines"] = s["receipt_lines"]
    return _render(f"receipt_{_suffix(b.basis)}.docx", ctx)


MINISTAY_ADDRESS_PLACEHOLDER = "[ADD MINISTAY ADDRESS]"


def visa_docx(b: BookingDocs, student_index: int, ministay: bool = False) -> bytes:
    """Ministay uses the same group letter; staff type the address into the Word file in Teams."""
    ctx = b.context()
    ctx["s"] = ctx["students"][student_index]
    if ministay:
        ctx["s"]["campus_address"] = MINISTAY_ADDRESS_PLACEHOLDER
    return _render(f"visa_{_suffix(b.basis)}.docx", ctx)


def docx_text(docx_bytes: bytes) -> str:
    """All visible text in a Word file (body and tables), for simple checks."""
    import docx as _docx
    d = _docx.Document(BytesIO(docx_bytes))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def to_pdf(docx_bytes: bytes) -> bytes:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "doc.docx"
        src.write_bytes(docx_bytes)
        subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "pdf",
                        "--outdir", tmp, str(src)], check=True, capture_output=True, timeout=180,
                       env={"HOME": tmp, "PATH": "/usr/bin:/bin:/usr/local/bin"})
        return (Path(tmp) / "doc.pdf").read_bytes()
